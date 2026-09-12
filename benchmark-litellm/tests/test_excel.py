from openpyxl import load_workbook

from bench.report.excel import build_workbook


def _records(label, n=5):
    return [
        {"api": "chat", "stream": True, "ok": i != 0, "reason": None if i != 0 else "content_mismatch",
         "latency_ms": 100 + i, "ttft_ms": 10, "overhead_ms": 5, "sched_lag_ms": 1,
         "t_sent_ms": i * 100, "t_done_ms": i * 100 + 100, "seed": f"{label}-{i}"}
        for i in range(n)
    ]


def test_build_workbook_has_expected_sheets(tmp_path):
    out_path = tmp_path / "report.xlsx"
    runs = {"baseline": _records("baseline"), "proxy_chat": _records("proxy_chat")}
    build_workbook(runs, out_path)
    wb = load_workbook(out_path)
    names = set(wb.sheetnames)
    assert "Summary" in names
    assert "By API x Stream" in names
    assert "Timeline" in names
    assert "Failures" in names
    assert "Raw-baseline" in names
    assert "Raw-proxy_chat" in names


def test_summary_sheet_has_a_row_per_run(tmp_path):
    out_path = tmp_path / "report.xlsx"
    runs = {"baseline": _records("baseline"), "proxy_chat": _records("proxy_chat")}
    build_workbook(runs, out_path)
    wb = load_workbook(out_path)
    ws = wb["Summary"]
    first_col = [c.value for c in ws["A"]]
    assert "baseline" in first_col
    assert "proxy_chat" in first_col
